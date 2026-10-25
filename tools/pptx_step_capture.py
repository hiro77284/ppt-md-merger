#!/usr/bin/env python
"""pptx のアニメーション段階をキャプチャする試作スクリプト（Windows 専用）。

docs/design_pptx_animation_capture.md の案 C（COM でスライドショーを駆動し、
ウィンドウをキャプチャする）の実現可能性を確認するための試作。本体パッケージ
（md_merge）には手を入れず、単独で動く。

確認したい点は 2 つ:
  1. ウィンドウ表示のスライドショーで GetClickCount() が期待どおりステップ数を返すか
  2. 1 ステップを安定してキャプチャできるか（差分収束待ちが効くか）

使い方:
  # PowerPoint を起動せずにタイムライン XML からステップ数を数える（確認 1 の事前チェック）
  python tools/pptx_step_capture.py FILE.pptx --inspect

  # PowerPoint を起動し、各スライドのステップ数だけ報告する（キャプチャしない）
  python tools/pptx_step_capture.py FILE.pptx --probe

  # スライド 3 の全ステップをキャプチャする
  python tools/pptx_step_capture.py FILE.pptx --slide 3 --outdir out

  # exportrange シェイプの範囲で切り出す（pptimgexport と同じ % 正規化クロップ）
  python tools/pptx_step_capture.py FILE.pptx --slide 3 --crop-shape exportrange

注意:
  * PowerPoint は不可視モードが使えないため、ウィンドウが画面に出る。全画面占有は
    しないが、最小化・セッションロック中は描画が止まるのでキャプチャできない。
  * 解像度は画面ピクセル依存（16:9 で 1920px なら約 144DPI 相当）。
"""

from __future__ import annotations

import argparse
import collections
import ctypes
import hashlib
import re
import sys
import time
from pathlib import Path

# ── PowerPoint COM 定数 ────────────────────────────────────────────────────
PP_SHOW_TYPE_WINDOW = 2          # ppShowTypeWindow: 全画面ではなくウィンドウで再生
PP_MANUAL_ADVANCE = 1            # ppSlideShowManualAdvance: 自動時間送りを止める
PP_SLIDE_SHOW_RUNNING = 1        # ppSlideShowRunning
PW_RENDERFULLCONTENT = 2         # PrintWindow のフラグ（DWM 合成内容を取る）


# ── 1. オフライン検査（PowerPoint 不要）───────────────────────────────────

def inspect_timing(pptx_path: Path) -> int:
    """タイムライン XML からクリックステップ数を数える（案 B の解析と同じ入口）。"""
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    print(f"{pptx_path}")
    print(f"  スライド寸法: {prs.slide_width} x {prs.slide_height} EMU")
    for idx, slide in enumerate(prs.slides, start=1):
        xml = slide._element.xml
        kinds = collections.Counter(re.findall(r'nodeType="(\w+)"', xml))
        presets = collections.Counter(re.findall(r'presetClass="(\w+)"', xml))
        clicks = kinds.get("clickEffect", 0)
        if not clicks and not presets:
            print(f"  slide {idx}: アニメーションなし")
            continue
        print(
            f"  slide {idx}: clickEffect={clicks} "
            f"withEffect={kinds.get('withEffect', 0)} "
            f"afterEffect={kinds.get('afterEffect', 0)} "
            f"presetClass={dict(presets)}"
        )
    return 0


# ── 2. 画面キャプチャ ──────────────────────────────────────────────────────

def _set_dpi_aware() -> None:
    """GetWindowRect が物理ピクセルを返すようにする（表示スケール 100% 以外で必要）。"""
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)  # PROCESS_PER_MONITOR_DPI_AWARE
    except Exception:
        try:
            ctypes.windll.user32.SetProcessDPIAware()
        except Exception:
            pass


def capture_window(hwnd: int, method: str = "printwindow"):
    """*hwnd* の内容を PIL.Image で返す。method は printwindow / bitblt。"""
    import win32con
    import win32gui
    import win32ui
    from PIL import Image

    left, top, right, bottom = win32gui.GetWindowRect(hwnd)
    width, height = right - left, bottom - top
    if width <= 0 or height <= 0:
        raise RuntimeError(f"ウィンドウサイズが不正です: {width}x{height}")

    hwnd_dc = win32gui.GetWindowDC(hwnd)
    mfc_dc = win32ui.CreateDCFromHandle(hwnd_dc)
    save_dc = mfc_dc.CreateCompatibleDC()
    bitmap = win32ui.CreateBitmap()
    bitmap.CreateCompatibleBitmap(mfc_dc, width, height)
    save_dc.SelectObject(bitmap)
    try:
        if method == "printwindow":
            ok = ctypes.windll.user32.PrintWindow(
                hwnd, save_dc.GetSafeHdc(), PW_RENDERFULLCONTENT
            )
            if not ok:
                raise RuntimeError("PrintWindow が失敗しました")
        else:
            save_dc.BitBlt((0, 0), (width, height), mfc_dc, (0, 0), win32con.SRCCOPY)
        info = bitmap.GetInfo()
        bits = bitmap.GetBitmapBits(True)
        img = Image.frombuffer(
            "RGB", (info["bmWidth"], info["bmHeight"]), bits, "raw", "BGRX", 0, 1
        )
    finally:
        save_dc.DeleteDC()
        mfc_dc.DeleteDC()
        win32gui.ReleaseDC(hwnd, hwnd_dc)
        win32gui.DeleteObject(bitmap.GetHandle())
    return img


def _is_blank(img) -> bool:
    """真っ黒（キャプチャ失敗の典型）かどうか。"""
    _lo, hi = img.convert("L").getextrema()
    return hi < 8


def client_box(hwnd: int) -> tuple[int, int, int, int]:
    """ウィンドウ画像内でのクライアント領域（タイトルバー・枠を除いた部分）の箱。"""
    import win32gui

    _cl, _ct, cr, cb = win32gui.GetClientRect(hwnd)
    sx, sy = win32gui.ClientToScreen(hwnd, (0, 0))
    wl, wt, _wr, _wb = win32gui.GetWindowRect(hwnd)
    ox, oy = sx - wl, sy - wt
    return (ox, oy, ox + cr, oy + cb)


def _band_trim(img, thresh: int = 12) -> tuple[int, int, int, int]:
    """四辺の暗い帯（レターボックス）を除いた箱を返す。Pillow だけで行・列平均を取る。"""
    from PIL import Image

    gray = img.convert("L")
    w, h = gray.size
    rows = list(gray.resize((1, h), Image.BOX).getdata())
    cols = list(gray.resize((w, 1), Image.BOX).getdata())
    top, bottom, left, right = 0, h, 0, w
    while top < bottom and rows[top] < thresh:
        top += 1
    while bottom > top and rows[bottom - 1] < thresh:
        bottom -= 1
    while left < right and cols[left] < thresh:
        left += 1
    while right > left and cols[right - 1] < thresh:
        right -= 1
    return (left, top, right, bottom)


def slide_box(img, aspect: float) -> tuple[int, int, int, int]:
    """キャプチャ画像の中のスライド領域を求める。

    ウィンドウ表示のスライドショーはスライドを黒帯で囲んで表示するので、まず暗い帯を
    実測で落とす。ブラウズ表示には右端に縦スクロールバーが付き、これは黒ではないため
    帯の検出では落ちない。そこで最後にスライドの縦横比へスナップして余剰幅を右から削る。
    """
    left, top, right, bottom = _band_trim(img)
    height = bottom - top
    want_w = round(height * aspect)
    if right - left > want_w:          # 右のスクロールバー等を落とす
        right = left + want_w
    elif right - left < want_w:        # 横が足りない（上下に余裕がある）場合
        want_h = round((right - left) / aspect)
        top += (height - want_h) // 2
        bottom = top + want_h
    return (left, top, right, bottom)


def capture_settled(hwnd: int, method: str, *, interval: float, timeout: float):
    """アニメーションが落ち着くまで撮り続け、2 枚連続で同一になったフレームを返す。

    完了イベントが無いため、フレーム差分の収束で代替する。
    戻り値: (image, 試行回数, 経過秒)
    """
    t0 = time.perf_counter()
    prev_digest = None
    prev_img = None
    attempts = 0
    while time.perf_counter() - t0 < timeout:
        img = capture_window(hwnd, method)
        attempts += 1
        digest = hashlib.blake2b(img.tobytes(), digest_size=16).digest()
        if digest == prev_digest:
            return prev_img, attempts, time.perf_counter() - t0
        prev_digest, prev_img = digest, img
        time.sleep(interval)
    return prev_img, attempts, time.perf_counter() - t0


# ── 3. クロップ（pptimgexport と同じ % 正規化）─────────────────────────────

def shape_rect_pct(pptx_path: Path, slide_no: int, shape_name: str):
    """シェイプ範囲をスライド寸法比（0..100）で返す。解像度非依存。"""
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    slide = list(prs.slides)[slide_no - 1]
    shape = next((s for s in slide.shapes if s.name == shape_name), None)
    if shape is None:
        raise SystemExit(f"シェイプが見つかりません: {shape_name!r}（slide {slide_no}）")
    sw, sh = prs.slide_width, prs.slide_height
    try:  # pptimgexport と同じ計算を共有できるなら使う
        from md_merge.pptimgexport import _to_ltrb_float

        return _to_ltrb_float(shape.left, shape.top, shape.width, shape.height, sw, sh)
    except Exception:
        return (
            shape.left / sw * 100,
            shape.top / sh * 100,
            (shape.left + shape.width) / sw * 100,
            (shape.top + shape.height) / sh * 100,
        )


def _slide_geometry(pptx_path: Path) -> tuple[float, float]:
    """(スライドの縦横比, スライド幅インチ) を返す。DPI 換算の表示に使う。"""
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    return prs.slide_width / prs.slide_height, prs.slide_width / 914400


def crop_pct(img, rect_pct):
    l, t, r, b = rect_pct
    w, h = img.size
    box = (
        int(round(w * l / 100)),
        int(round(h * t / 100)),
        int(round(w * r / 100)),
        int(round(h * b / 100)),
    )
    return img.crop(box)


# ── 4. スライドショー駆動 ─────────────────────────────────────────────────

def run_show(args) -> int:
    import pythoncom
    import win32com.client as win32

    pptx_path = Path(args.pptx).resolve()
    if not pptx_path.exists():
        print(f"ファイルがありません: {pptx_path}", file=sys.stderr)
        return 3

    _set_dpi_aware()
    pythoncom.CoInitialize()
    app = win32.Dispatch("PowerPoint.Application")
    app.Visible = True   # PowerPoint は不可視モードが使えない

    pres = app.Presentations.Open(str(pptx_path), ReadOnly=True, WithWindow=True)
    show_window = None
    try:
        settings = pres.SlideShowSettings
        settings.AdvanceMode = PP_MANUAL_ADVANCE
        if not args.fullscreen:
            settings.ShowType = PP_SHOW_TYPE_WINDOW
        for prop, value in (("ShowPresenterView", False), ("ShowWithAnimation", True)):
            try:
                setattr(settings, prop, value)
            except Exception as e:  # バージョンによって無い／設定できない
                print(f"  (情報) {prop} を設定できませんでした: {e}")

        windows_before = list_ppt_windows()
        show_window = settings.Run()
        view = show_window.View
        time.sleep(args.startup_wait)
        print(
            f"スライドショー開始: state={view.State} "
            f"size={show_window.Width:.0f}x{show_window.Height:.0f}pt"
        )

        slide_count = pres.Slides.Count
        targets = [args.slide] if args.slide else list(range(1, slide_count + 1))

        # ── 各スライドのステップ数を報告（確認 1）──────────────────────
        click_counts = {}
        for n in targets:
            view.GotoSlide(n)
            time.sleep(args.settle_interval)
            try:
                clicks = view.GetClickCount()
            except Exception as e:
                print(f"  slide {n}: GetClickCount() 失敗: {e}")
                clicks = 0
            click_counts[n] = clicks
            print(f"  slide {n}: GetClickCount() = {clicks}")

        if args.probe:
            return 0

        # ── キャプチャ（確認 2）────────────────────────────────────────
        hwnd = _show_hwnd(show_window, windows_before)
        print(f"  キャプチャ対象 HWND = {hwnd}  method={args.method}")
        outdir = Path(args.outdir).resolve()
        outdir.mkdir(parents=True, exist_ok=True)

        rect_pct = None
        if args.crop_shape:
            rect_pct = shape_rect_pct(pptx_path, targets[0], args.crop_shape)
            print(
                f"  crop {args.crop_shape}: l={rect_pct[0]:.2f}% t={rect_pct[1]:.2f}% "
                f"r={rect_pct[2]:.2f}% b={rect_pct[3]:.2f}%"
            )

        slide_aspect, slide_in_w = _slide_geometry(pptx_path)
        trim_box = None if args.no_trim else client_box(hwnd)

        results = []
        for n in targets:
            clicks = click_counts[n]
            for step in range(clicks + 1):   # step 0 = 最初のクリック前の状態
                _goto(view, n, step, use_next=args.use_next)
                img, attempts, elapsed = capture_settled(
                    hwnd,
                    args.method,
                    interval=args.settle_interval,
                    timeout=args.settle_timeout,
                )
                if img is None:
                    print(f"  slide {n} step {step}: キャプチャできませんでした")
                    continue
                blank = _is_blank(img)
                if trim_box is not None:
                    # ウィンドウ枠・タイトルバーを除き、さらに黒帯を取ってスライド領域だけにする
                    client_img = img.crop(trim_box)
                    box = slide_box(client_img, slide_aspect)
                    img = client_img.crop(box)
                    if step == 0 and n == targets[0]:
                        dpi = img.size[0] / slide_in_w
                        print(
                            f"  スライド領域: client 内 {box} -> {img.size[0]}x{img.size[1]}px "
                            f"（{dpi:.0f} DPI 相当。pptimgexport の既定は 150）"
                        )
                if rect_pct:
                    img = crop_pct(img, rect_pct)
                out = outdir / f"slide{n}_step{step}.png"
                img.save(out)
                results.append((n, step, img.size, attempts, elapsed, blank))
                warn = "  ← 真っ黒（キャプチャ失敗の可能性）" if blank else ""
                print(
                    f"  slide {n} step {step}: {img.size[0]}x{img.size[1]} "
                    f"frames={attempts} settle={elapsed:.2f}s{warn} -> {out.name}"
                )

        _print_summary(results)
        return 0
    finally:
        try:
            if show_window is not None and show_window.View.State == PP_SLIDE_SHOW_RUNNING:
                show_window.View.Exit()
        except Exception:
            pass
        try:
            pres.Close()
        except Exception:
            pass
        if not args.keep_open:
            try:
                if app.Presentations.Count == 0:
                    app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def list_ppt_windows() -> dict[int, tuple[int, int]]:
    """可視の PowerPoint フレームウィンドウ（class=PPTFrameClass）の {hwnd: (w, h)}。

    ウィンドウ表示のスライドショーは、PowerPoint 本体と同じ ``PPTFrameClass`` の
    別ウィンドウとして出る（タイトルは「PowerPoint スライド ショー - …」だが
    ローカライズされるため当てにしない）。
    """
    import win32gui

    found: dict[int, tuple[int, int]] = {}

    def cb(hwnd, _):
        if not win32gui.IsWindowVisible(hwnd):
            return
        if win32gui.GetClassName(hwnd) != "PPTFrameClass":
            return
        left, top, right, bottom = win32gui.GetWindowRect(hwnd)
        found[hwnd] = (right - left, bottom - top)

    win32gui.EnumWindows(cb, None)
    return found


def _show_hwnd(show_window, before: dict[int, tuple[int, int]]) -> int:
    """スライドショーウィンドウの HWND を特定する。

    1. ``SlideShowWindow.HWND``（遅延バインディングではメソッドとして見えるので両方試す。
       PowerPoint 2016 以降では取得できないことがある）
    2. スライドショー起動の前後で増えた ``PPTFrameClass`` ウィンドウ（ローカライズ非依存）
    3. COM が報告するウィンドウサイズ（pt）に一致するウィンドウ
    """
    raw = getattr(show_window, "HWND", None)
    if callable(raw):
        try:
            raw = raw()
        except Exception:
            raw = None
    if raw is not None:
        try:
            return int(raw)
        except (TypeError, ValueError):
            pass

    after = list_ppt_windows()
    new = [h for h in after if h not in before]
    if new:
        hwnd = max(new, key=lambda h: after[h][0] * after[h][1])
        print(f"  (情報) 起動前後の差分からスライドショーウィンドウを特定: hwnd={hwnd}")
        return hwnd

    # 1pt = 1/72 inch、既定 96 DPI なら 4/3 px
    try:
        want_w = round(float(show_window.Width) * 4 / 3)
        want_h = round(float(show_window.Height) * 4 / 3)
        for hwnd, (w, h) in after.items():
            if abs(w - want_w) <= 8 and abs(h - want_h) <= 8:
                print(f"  (情報) サイズ一致でスライドショーウィンドウを特定: hwnd={hwnd}")
                return hwnd
    except Exception:
        pass
    raise RuntimeError("スライドショーウィンドウの HWND を取得できませんでした")


def _goto(view, slide_no: int, step: int, *, use_next: bool) -> None:
    """スライド *slide_no* のステップ *step* の状態へ移動する。

    既定は GotoClick（任意ステップへ決定的に跳べる。やり直しも効く）。
    --use-next 指定時は GotoSlide + Next() の繰り返しで、実際のクリック送りを再現する。
    """
    if not use_next:
        view.GotoSlide(slide_no)
        if step:
            view.GotoClick(step)
        return
    view.GotoSlide(slide_no)
    for _ in range(step):
        view.Next()


def _print_summary(results) -> None:
    if not results:
        print("\n結果なし")
        return
    blanks = [r for r in results if r[5]]
    frames = [r[3] for r in results]
    settle = [r[4] for r in results]
    sizes = {r[2] for r in results}
    print("\n--- まとめ ---")
    print(f"  キャプチャ枚数: {len(results)}")
    print(f"  真っ黒だった枚数: {len(blanks)}")
    print(f"  収束までのフレーム数: min={min(frames)} max={max(frames)}")
    print(f"  収束までの時間: min={min(settle):.2f}s max={max(settle):.2f}s")
    print(f"  画像サイズ: {', '.join(f'{w}x{h}' for w, h in sorted(sizes))}")


# ── 5. CLI ────────────────────────────────────────────────────────────────

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("pptx", help="対象 pptx")
    ap.add_argument(
        "--inspect", action="store_true",
        help="PowerPoint を起動せず、タイムライン XML からステップ数を数えるだけ",
    )
    ap.add_argument(
        "--probe", action="store_true",
        help="スライドショーを起動して GetClickCount() を報告するだけ（キャプチャしない）",
    )
    ap.add_argument("--slide", type=int, help="対象スライド番号（省略時は全スライド）")
    ap.add_argument("--outdir", default="step_capture_out", help="PNG の出力先")
    ap.add_argument("--crop-shape", help="このシェイプの範囲で切り出す（例 exportrange）")
    ap.add_argument(
        "--method", choices=("printwindow", "bitblt"), default="printwindow",
        help="キャプチャ方法。printwindow が黒くなるなら bitblt（前面必須）を試す",
    )
    ap.add_argument(
        "--use-next", action="store_true",
        help="GotoClick ではなく Next() の連打でステップを進める",
    )
    ap.add_argument(
        "--fullscreen", action="store_true",
        help="ウィンドウ表示をやめて全画面で再生する（画面を占有する）",
    )
    ap.add_argument(
        "--no-trim", action="store_true",
        help="ウィンドウ枠と黒帯を除かず、ウィンドウ全体をそのまま保存する（デバッグ用）",
    )
    ap.add_argument("--startup-wait", type=float, default=1.5, help="スライドショー起動後の待ち秒数")
    ap.add_argument("--settle-interval", type=float, default=0.12, help="差分収束判定の撮影間隔（秒）")
    ap.add_argument("--settle-timeout", type=float, default=3.0, help="差分収束の打ち切り秒数")
    ap.add_argument("--keep-open", action="store_true", help="終了時に PowerPoint を終了しない")
    args = ap.parse_args(argv)

    if args.inspect:
        return inspect_timing(Path(args.pptx).resolve())
    if sys.platform != "win32":
        print("このスクリプトは Windows 専用です（PowerPoint COM を使用）", file=sys.stderr)
        return 2
    return run_show(args)


if __name__ == "__main__":
    sys.exit(main())
