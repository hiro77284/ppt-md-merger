#!/usr/bin/env python3
"""pptx ノートの export-image-mode を全走査し、不正値・書きかけを検出する（read-only）。

判定は pptimgexport と同じ ``parse_note_image_directives`` / ``_KNOWN_MODES`` を使うので、
実行時の警告と一致する。編集後の確認（typo の再混入チェック）に使える。

分類:
  - INVALID   … mode が非空だが不正値（typo 等）。要修正
  - PARTIAL   … mode 未指定だが他の export-image-* に値がある（書きかけの不備）
  - (参考)     … 全項目が空の IMAGE ブロック＝プレースホルダ。件数のみ集計

使い方:
    python tools/scan_export_mode.py [ROOT]      # 既定 ROOT=カレントディレクトリ
    python tools/scan_export_mode.py . --all      # build/_bak も含める

終了コード: INVALID か PARTIAL が 1 件以上あれば 1、無ければ 0。
依存: python-pptx（＋ md_merge パッケージ）。
"""
from __future__ import annotations

import sys
from pathlib import Path

# 未インストールでも動くよう同梱 src を優先的に見えるようにする
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from pptx import Presentation  # noqa: E402

from md_merge._note_parser import parse_note_image_directives  # noqa: E402
from md_merge.pptimgexport import _KNOWN_MODES  # noqa: E402


def main() -> int:
    argv = [a for a in sys.argv[1:] if not a.startswith("-")]
    include_all = "--all" in sys.argv
    root = Path(argv[0]) if argv else Path(".")

    def skip(p: Path) -> bool:
        if p.name.startswith("~$"):
            return True
        if include_all:
            return False
        return "_pptx_latex_bak" in p.parts or "build" in p.parts

    targets = [p for p in sorted(root.rglob("*.pptx")) if not skip(p)]

    invalid: list[tuple[Path, int, str]] = []
    partial: list[tuple[Path, int]] = []
    empty = 0
    errors: list[tuple[Path, str]] = []

    for p in targets:
        try:
            prs = Presentation(str(p))
        except Exception as e:  # noqa: BLE001  ロック等
            errors.append((p, str(e)))
            continue
        for idx, slide in enumerate(prs.slides, start=1):
            if not slide.has_notes_slide:
                continue
            note = slide.notes_slide.notes_text_frame.text or ""
            for meta in parse_note_image_directives(note):
                mode = (meta.get("export-image-mode") or "").strip()
                if mode in _KNOWN_MODES:
                    continue
                rel = p.relative_to(root)
                if mode:
                    invalid.append((rel, idx, mode))
                elif (meta.get("export-image") or "").strip() or (meta.get("export-image-target") or "").strip():
                    partial.append((rel, idx))
                else:
                    empty += 1

    print(f"走査 {len(targets)} pptx / 有効値: {', '.join(sorted(_KNOWN_MODES))}\n")

    print(f"=== INVALID（不正値・typo）: {len(invalid)} 件 ===")
    for rel, idx, val in invalid:
        print(f"  {rel}  slide {idx}: {val!r}")
    if not invalid:
        print("  (なし)")

    print(f"\n=== PARTIAL（mode 未指定・他項目あり）: {len(partial)} 件 ===")
    for rel, idx in partial:
        print(f"  {rel}  slide {idx}")
    if not partial:
        print("  (なし)")

    print(f"\n参考: 全項目が空の IMAGE ブロック（プレースホルダ）: {empty} 件（無警告でスキップされる）")

    if errors:
        print(f"\n=== 開けなかったファイル: {len(errors)} 件（ロック等）===")
        for p, e in errors:
            print(f"  {p.name}: {e}")

    return 1 if (invalid or partial) else 0


if __name__ == "__main__":
    sys.exit(main())
