"""pptyamlexport — PPTX ノートの YAML_BLOCK から図版定義 YAML を書き出す。

ノート本文は投入先 YAML そのものをテンプレートとして書き、slot の座標には
``auto`` を置く。``auto`` はスライド上の名前付きシェイプの実位置から
``baseframe`` 内の相対値（0〜1）として算出して埋める。トップレベルの
``regions`` も slots と同じ規則で解決する。

ノート本文はレシピ YAML と同じ規則（``_include`` / ``_constants`` /
``__CONSTANT__``）で読み込むので、共通の slot 定義を外部ファイル
に切り出して複数スライドから取り込める。

pptmdexport とは扱うデータが根本的に異なる（散文 MD ではなく構造化データ）
ため別サブコマンドとして独立させている。共通するのは cropsrc の解決と
スライド走査だけ。
"""

from __future__ import annotations

import argparse
import json
import logging
from pathlib import Path

import yaml
from pptx import Presentation

from md_merge._filters import replace_slide_placeholders
from md_merge._note_parser import parse_note_image_directives, parse_note_yaml_blocks
from md_merge._cropsrc import cropsrc_entries, resolve_cropsrc_path
from md_merge._shapes import find_shape as _find_shape
from md_merge._output import EXIT_BAD_INPUT, EXIT_FAILURE, EXIT_NOT_FOUND, EXIT_OK, emit, setup_logging
from md_merge._preprocess import cropsrc_read_path
from md_merge.merge._recipe import (
    load_yaml_text,
    load_yaml,
    resolve_input_path,
    resolve_workdir,
    setup_recipe_file_logging,
)

_AUTO = "auto"

# type: module の定義は slots を持たない（regions だけの部品定義）。
_MODULE_TYPE = "module"

# kind: reading の定義は slots を任意にする（slot を持たない読み物定義）。
_READING_KIND = "reading"

# baseframe の予約値。スライド全体を座標の基準にする。export-image-mode の
# 同名の値（スライド全体を切り出すモード）と対になっている。
_SLIDE_BASEFRAME = "slide"

# baseframe 内相対値として算出できる slot キー。
_GEOM_KEYS = ("x", "y", "width", "height")

# 出力 YAML には残さない、ノート側の指示専用キー。
_AUTHORING_KEYS = ("export-note", "baseframe")

_ALIGN_H = {"left", "center", "right"}
_ALIGN_V = {"top", "middle", "bottom"}

# 相対値の丸め桁数。EMU 由来の除算で出る 0.30000000000000004 のような
# 浮動小数点ノイズを落とす。1e-4 は PNG 数千 px でもサブピクセル精度。
_ROUND = 4


class _SlideRect:
    """スライド全体を baseframe と同じ left/top/width/height で見せる矩形。

    ``export-image-mode: slide`` で切り出した PNG はスライド全体そのものなので、
    座標の基準もスライド矩形になる。シェイプと同じ属性を持たせて
    ``_rel_geometry`` をそのまま使えるようにしている。
    """

    def __init__(self, prs):
        self.left = 0
        self.top = 0
        self.width = prs.slide_width
        self.height = prs.slide_height


def _rel_geometry(shape, base) -> dict[str, float]:
    """*shape* の位置を *base* の矩形内の相対値（0〜1）に変換する。"""
    return {
        "x": (shape.left - base.left) / base.width,
        "y": (shape.top - base.top) / base.height,
        "width": shape.width / base.width,
        "height": shape.height / base.height,
    }


def _is_auto(slot: dict, key: str) -> bool:
    """*key* が auto 扱いか（明示の ``auto`` またはキー自体の省略）を判定する。

    座標キーの省略は ``auto`` の省略記法とみなす。4 つとも書かずに済むのが
    ふつうの使い方で、``auto`` を 4 行並べるのは冗長なため。
    """
    if key not in slot:
        return True
    return str(slot[key]).strip().lower() == _AUTO


def _set_after(doc: dict, key: str, value, after: tuple[str, ...]) -> None:
    """*after* のいずれか最初に見つかったキーの直後に *key* を挿入する。

    dict は挿入順を保つので、末尾に追記されると投入先 YAML のキー順が崩れる。
    どのアンカーも無ければ末尾に置く。
    """
    if key in doc:
        doc[key] = value
        return
    anchor = next((k for k in after if k in doc), None)
    if anchor is None:
        doc[key] = value
        return
    rebuilt: dict = {}
    for k, v in doc.items():
        rebuilt[k] = v
        if k == anchor:
            rebuilt[key] = value
    doc.clear()
    doc.update(rebuilt)


def _slide_image_block(note_text: str) -> dict[str, str] | None:
    """スライド全体を切り出す IMAGE ブロックを返す（無ければ None）。"""
    for meta in parse_note_image_directives(note_text):
        if meta.get("export-image-mode", "").strip() == _SLIDE_BASEFRAME:
            return meta
    return None


def _image_block_for(note_text: str, baseframe: str) -> dict[str, str] | None:
    """baseframe と同じ範囲を切り出す IMAGE ブロックを返す（無ければ None）。"""
    if baseframe == _SLIDE_BASEFRAME:
        return _slide_image_block(note_text)
    for meta in parse_note_image_directives(note_text):
        if meta.get("export-image-target", "").strip() == baseframe:
            return meta
    return None


def _base_desc(baseframe: str) -> str:
    """診断文に埋める座標基準の呼び名。"""
    if baseframe == _SLIDE_BASEFRAME:
        return "スライド全体"
    return f"baseframe '{baseframe}'"


def _check_image_consistency(doc: dict, note_text: str, baseframe: str, where: str) -> None:
    """image: と、baseframe を切り出す IMAGE ブロックの整合を検査する。

    slot 座標は baseframe の矩形を基準に算出するため、PNG の実際の範囲が
    baseframe とズレると座標が丸ごと狂う。ズレの原因になる
    export-image-inset を検出して警告し、image: 未指定なら IMAGE ブロックの
    出力名で補完する。
    """
    desc = _base_desc(baseframe)
    meta = _image_block_for(note_text, baseframe)
    if meta is None:
        logging.warning(
            "%s: %sを切り出す IMAGE ブロックが同じスライドに見つかりません。"
            "image: と座標基準の整合を検証できません",
            where, desc,
        )
        return

    inset = (meta.get("export-image-inset") or "").strip()
    if inset and inset not in ("0", "+0", "-0"):
        logging.warning(
            "%s: %sの IMAGE ブロックに export-image-inset: %s が指定されています。"
            "PNG の範囲が座標基準より %s px 分ずれるため slot 座標が合いません"
            "（inset なしの切り出しを使ってください）",
            where, desc, inset, inset,
        )

    exported = Path((meta.get("export-image") or "").strip()).name
    if not exported:
        return
    current = str(doc.get("image") or "").strip()
    if not current:
        _set_after(doc, "image", exported, ("title", "id"))
        logging.debug("%s: image: を IMAGE ブロックの出力名で補完しました: %s", where, exported)
    elif Path(current).name != exported:
        logging.warning(
            "%s: image: %s が%sの切り出し先 %s と一致しません",
            where, current, desc, exported,
        )


def _warn_out_of_range(geom: dict, ident: str, where: str) -> None:
    """座標が baseframe の外にはみ出していれば警告する（値は補正しない）。"""
    vals: dict[str, float] = {}
    for k in _GEOM_KEYS:
        v = geom.get(k)
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            vals[k] = float(v)
    for k, v in vals.items():
        if not (0.0 <= v <= 1.0):
            logging.warning("%s: %s の %s が 0〜1 の範囲外です: %s", where, ident, k, v)
    tol = 10 ** -_ROUND
    if "x" in vals and "width" in vals and vals["x"] + vals["width"] > 1.0 + tol:
        logging.warning(
            "%s: %s が baseframe の右端をはみ出します (x+width=%.4f)",
            where, ident, vals["x"] + vals["width"],
        )
    if "y" in vals and "height" in vals and vals["y"] + vals["height"] > 1.0 + tol:
        logging.warning(
            "%s: %s が baseframe の下端をはみ出します (y+height=%.4f)",
            where, ident, vals["y"] + vals["height"],
        )


def _warn_align(slot: dict, ident: str, where: str) -> None:
    """text_align / vertical_align の語彙を検査する。"""
    h = slot.get("text_align")
    if h is not None and str(h) not in _ALIGN_H:
        logging.warning(
            "%s: %s の text_align が不正です: %r（有効値: %s）",
            where, ident, h, ", ".join(sorted(_ALIGN_H)),
        )
    v = slot.get("vertical_align")
    if v is not None and str(v) not in _ALIGN_V:
        logging.warning(
            "%s: %s の vertical_align が不正です: %r（有効値: %s）",
            where, ident, v, ", ".join(sorted(_ALIGN_V)),
        )


def _lookup_shape(slide, name: str, where: str, subject: str):
    """座標の基準に使うシェイプを引く。slots と regions で共有する。

    *subject* は診断文に埋める主語（「slot 'slotA' と同名のシェイプ」など）。

    Returns:
        (shape, エラー理由)。見つからなければ shape が None。
    """
    shape, in_group = _find_shape(slide, name)
    if shape is None:
        if in_group:
            return None, (
                f"{where}: {subject}はグループ内にあります。"
                "グループ内シェイプの座標は親の子座標系で表されるため使えません"
                "（グループを解除してください）"
            )
        return None, f"{where}: {subject}が見つかりません"

    if getattr(shape, "rotation", 0):
        logging.warning(
            "%s: %sは %s 度回転しています。"
            "境界ボックスから算出するため回転は反映されません",
            where, subject, shape.rotation,
        )
    return shape, None


def _auto_values(
    geom: dict, name: str, slide, base, where: str, ident: str
) -> tuple[dict[str, float] | None, str | None]:
    """*geom* の auto（省略を含む）座標を、*name* と同名のシェイプから算出する。

    Returns:
        (算出した {キー: 相対値}, エラー理由)。エラーのとき第 1 要素は None。
    """
    auto_keys = [k for k in _GEOM_KEYS if _is_auto(geom, k)]
    if not auto_keys:
        return {}, None

    shape, problem = _lookup_shape(slide, name, where, f"{ident} と同名のシェイプ")
    if problem:
        return None, problem
    rel = _rel_geometry(shape, base)
    return {k: round(rel[k], _ROUND) for k in auto_keys}, None


def _resolve_items(items, plural: str, singular: str, slide, base, where: str) -> list[str]:
    """``id`` と同名のシェイプから auto 座標を解決する。エラー理由のリストを返す。

    slots と regions は「id ＋ 座標 ＋ 投入先向けの付加キー」という同じ形なので
    解決処理を共有する。``label`` / ``snap`` など座標以外のキーは解釈せず素通し
    する。
    """
    problems: list[str] = []
    if not isinstance(items, list):
        problems.append(f"{where}: {plural} がリストではありません")
        return problems

    seen: set[str] = set()
    for i, item in enumerate(items):
        if not isinstance(item, dict):
            problems.append(f"{where}: {plural}[{i}] がマッピングではありません")
            continue

        item_id = str(item.get("id") or "").strip()
        if not item_id:
            problems.append(f"{where}: {plural}[{i}] に id がありません")
            continue
        if item_id in seen:
            problems.append(f"{where}: {singular} id が重複しています: {item_id}")
            continue
        seen.add(item_id)

        ident = f"{singular} '{item_id}'"
        values, problem = _auto_values(item, item_id, slide, base, where, ident)
        if problem:
            problems.append(problem)
            continue
        for k, value in values.items():
            if k in item:
                item[k] = value
                continue
            # 省略されていたキーは末尾に足すとキー順が崩れるので、直前の
            # 座標キー（無ければ id）の後ろに差し込んで x→y→width→height
            # の並びを保つ。values は _GEOM_KEYS 順なので、この時点で
            # 先行するキーは既に挿入済み。
            preceding = _GEOM_KEYS[: _GEOM_KEYS.index(k)]
            _set_after(item, k, value, tuple(reversed(preceding)) + ("id",))

        _warn_out_of_range(item, ident, where)
        _warn_align(item, ident, where)

    return problems


def _is_module(doc: dict) -> bool:
    """``type: module`` の定義か（slots を持たない部品定義か）を判定する。"""
    return str(doc.get("type") or "").strip().lower() == _MODULE_TYPE


def _is_reading(doc: dict) -> bool:
    """``kind: reading`` の定義か（slots を持たなくてよい読み物定義か）を判定する。"""
    return str(doc.get("kind") or "").strip().lower() == _READING_KIND


def _resolve_slots(doc: dict, slide, base, where: str) -> list[str]:
    """slot の auto 座標を名前付きシェイプから解決する。

    slots は通常必須だが、``type: module`` と ``kind: reading`` の定義では
    任意にする（regions だけの部品定義、slot を持たない読み物定義を書ける
    ようにするため）。省略を許すだけで、書かれていれば扱いは同じ（auto も
    従来どおり解決する）。
    """
    if doc.get("slots") is None and (_is_module(doc) or _is_reading(doc)):
        return []
    return _resolve_items(doc.get("slots"), "slots", "slot", slide, base, where)


def _resolve_regions(doc: dict, slide, base, where: str) -> list[str]:
    """region の auto 座標を名前付きシェイプから解決する（regions は任意）。"""
    if doc.get("regions") is None:
        return []
    return _resolve_items(doc.get("regions"), "regions", "region", slide, base, where)


def _warn_legacy_layout(doc: dict, where: str) -> None:
    """廃止した samples[].layout が残っていれば警告する。

    値はそのまま出力へ通す（投入先が無視する想定）。座標の自動算出は行わない。
    """
    samples = doc.get("samples")
    for i, sample in enumerate(samples if isinstance(samples, list) else []):
        if not isinstance(sample, dict) or "layout" not in sample:
            continue
        key = str(sample.get("key") or "").strip() or f"[{i}]"
        logging.warning(
            "%s: samples '%s' に layout が書かれていますが layout は廃止されました"
            "（トップレベルの regions を使ってください）。座標は算出せずそのまま出力します",
            where, key,
        )


class _IndentedDumper(yaml.SafeDumper):
    """シーケンスを親キーの下にインデントして出力する Dumper。

    PyYAML の既定は ``slots:`` の直下に ``- id:`` を字下げなしで並べる。投入先
    YAML は人が読み書きする定義ファイルなので、手書きの体裁（2 スペース字下げ）
    に揃える。
    """

    def increase_indent(self, flow=False, indentless=False):
        return super().increase_indent(flow, False)


def _dump_doc(doc: dict) -> str:
    return yaml.dump(
        doc, Dumper=_IndentedDumper,
        allow_unicode=True, default_flow_style=False, sort_keys=False,
    )


def _load_note_yaml(body: str, base_dir: Path):
    """ノートの YAML ブロック本文を、レシピ YAML と同じ前処理を通して読み込む。

    ``_include`` / ``_constants`` / ``__CONSTANT__``（旧書式の ``!include`` /
    ``!constants`` も）はレシピ YAML とまったく同じ規則で展開する。相対パスの基準は **元 PPTX と同じディレクトリ**で、``export-note``
    / ``export-image`` の解決基準と揃えてある。

    展開はエクスポート時に行われ、出力 YAML には展開後の内容が書き出される
    （投入先は ``_include`` を解釈しないため）。
    """
    return load_yaml_text(body, base_dir)


def _needs_auto(doc: dict) -> bool:
    """slots / regions のいずれかの座標が auto か（省略されているか）を判定する。"""
    for section in ("slots", "regions"):
        items = doc.get(section)
        for item in (items if isinstance(items, list) else []):
            if not isinstance(item, dict):
                continue
            for k in _GEOM_KEYS:
                if _is_auto(item, k):
                    return True
    return False


def run(args: argparse.Namespace) -> int:
    setup_logging(args.log_level)

    yaml_path = resolve_input_path(args)
    if yaml_path is None:
        logging.error("No input file specified")
        return EXIT_BAD_INPUT
    if not yaml_path.exists():
        logging.error("Input file not found: %s", yaml_path)
        return EXIT_NOT_FOUND

    recipe = load_yaml(yaml_path)
    entries = cropsrc_entries(recipe)
    output_section = recipe.get("output") or {}
    cli_workdir = Path(args.workdir).resolve() if getattr(args, "workdir", None) else None
    workdir = resolve_workdir(recipe, yaml_path, cli_workdir)
    setup_recipe_file_logging(recipe, yaml_path, workdir)
    force = getattr(args, "force", False) or bool(output_section.get("force"))
    dry_run = getattr(args, "dry_run", False)
    json_mode = getattr(args, "json", False)

    if not entries:
        logging.error("input.cropsrc が指定されていません: %s", yaml_path)
        return EXIT_BAD_INPUT

    # ── 全エントリのパス解決と存在確認（fail fast）────────────────────────
    resolved_pptx_paths: list[tuple[str, Path]] = []
    for i, entry in enumerate(entries):
        label = f"input.cropsrc[{i}]" if len(entries) > 1 else "input.cropsrc"
        pptx_val = entry.get("pptx")
        if not pptx_val:
            logging.error("%s.pptx が指定されていません: %s", label, yaml_path)
            return EXIT_BAD_INPUT
        pptx_path = resolve_cropsrc_path(str(pptx_val), yaml_path, workdir)
        if not pptx_path.exists():
            logging.error("PPTX ファイルが見つかりません: %s", pptx_path)
            return EXIT_NOT_FOUND
        resolved_pptx_paths.append((str(pptx_val), pptx_path))

    # ── エントリごとに処理 ────────────────────────────────────────────────
    errors = 0
    failures: list[str] = []
    json_entries: list[dict] = []

    for pptx_val, pptx_path in resolved_pptx_paths:
        # --preprocess 時は前処理済み pptx から読む（出力パスとプレースホルダー
        # の解決は元の pptx_path のまま維持する）。
        read_pptx_path = cropsrc_read_path(args, pptx_path)
        try:
            prs = Presentation(str(read_pptx_path))
        except Exception as e:
            logging.error("PPTX ファイルを開けません: %s: %s", read_pptx_path, e)
            errors += 1
            failures.append(f"PPTX を開けない: {read_pptx_path}: {e}")
            continue

        tasks: list[dict] = []
        for slide_idx, slide in enumerate(prs.slides, start=1):
            note_text = ""
            if slide.has_notes_slide:
                tf = slide.notes_slide.notes_text_frame
                note_text = (tf.text or "") if tf else ""

            title_shape = slide.shapes.title
            slide_title = title_shape.text if title_shape is not None else ""
            note_text = replace_slide_placeholders(note_text, pptx_path, slide_title, slide_idx)

            for block_idx, body in enumerate(parse_note_yaml_blocks(note_text)):
                where = f"スライド {slide_idx}"
                if block_idx:
                    where += f" (YAML ブロック {block_idx + 1})"

                if not body.strip():
                    continue  # 空の YAML ブロックはプレースホルダとみなし黙ってスキップ

                try:
                    doc = _load_note_yaml(body, pptx_path.parent)
                except yaml.YAMLError as e:
                    logging.error("%s: YAML として解析できません: %s", where, e)
                    errors += 1
                    failures.append(f"{where}: YAML 解析失敗: {e}")
                    continue
                except (OSError, ValueError) as e:
                    # _include 先が無い / 循環している。前処理で潰れるとどの
                    # スライドが原因か分からなくなるので、ここで場所を添える。
                    logging.error("%s: YAML の前処理に失敗しました: %s", where, e)
                    errors += 1
                    failures.append(f"{where}: _include 展開失敗: {e}")
                    continue

                if not isinstance(doc, dict):
                    logging.error("%s: YAML のトップレベルがマッピングではありません", where)
                    errors += 1
                    failures.append(f"{where}: トップレベルがマッピングでない")
                    continue

                doc_id = str(doc.get("id") or "").strip()
                if not doc_id:
                    logging.error("%s: id が指定されていません", where)
                    errors += 1
                    failures.append(f"{where}: id 未指定")
                    continue

                baseframe = str(doc.get("baseframe") or "").strip()
                base = None
                if baseframe.lower() == _SLIDE_BASEFRAME:
                    # 予約値。スライド全体が座標の基準。
                    baseframe = _SLIDE_BASEFRAME
                    base = _SlideRect(prs)
                elif baseframe:
                    base, in_group = _find_shape(slide, baseframe)
                    if base is None:
                        detail = (
                            f"{where}: baseframe '{baseframe}' のシェイプはグループ内にあります"
                            "（グループを解除してください）"
                            if in_group
                            else f"{where}: baseframe '{baseframe}' と同名のシェイプが見つかりません"
                        )
                        logging.error("%s", detail)
                        errors += 1
                        failures.append(detail)
                        continue
                elif _needs_auto(doc):
                    # baseframe 省略時、同じスライドの IMAGE ブロックがスライド全体を
                    # 切り出しているなら PNG の範囲はスライドそのものなので、基準も
                    # スライド矩形に決まる。クロップ画像（named_shape /
                    # slide_with_rectpct）では基準が定まらないので従来どおりエラー。
                    if _slide_image_block(note_text) is not None:
                        baseframe = _SLIDE_BASEFRAME
                        base = _SlideRect(prs)
                        logging.debug(
                            "%s: baseframe 省略、IMAGE ブロックが export-image-mode: %s のため"
                            "スライド全体を座標の基準にします",
                            where, _SLIDE_BASEFRAME,
                        )
                    else:
                        detail = f"{where}: slot に auto（または省略された）座標がありますが baseframe が指定されていません"
                        logging.error("%s", detail)
                        errors += 1
                        failures.append(detail)
                        continue

                if base is not None:
                    _check_image_consistency(doc, note_text, baseframe, where)

                problems = _resolve_slots(doc, slide, base, where)
                problems += _resolve_regions(doc, slide, base, where)
                _warn_legacy_layout(doc, where)
                if problems:
                    for p in problems:
                        logging.error("%s", p)
                    errors += len(problems)
                    failures.extend(problems)
                    continue

                export_val = str(doc.get("export-note") or "").strip() or f"{doc_id}.yaml"
                if not export_val.lower().endswith((".yaml", ".yml")):
                    logging.warning(
                        "%s: export-note のファイル名が .yaml ではありません: %s", where, export_val
                    )

                # ノート側の指示専用キーは投入先 YAML には残さない。
                for k in _AUTHORING_KEYS:
                    doc.pop(k, None)

                export_path = Path(export_val)
                out_path = export_path if export_path.is_absolute() else (pptx_path.parent / export_path).resolve()

                tasks.append({
                    "slide": slide_idx,
                    "id": doc_id,
                    "baseframe": baseframe,
                    "slots": len(doc.get("slots") or []),
                    "output": out_path,
                    "doc": doc,
                })

        if not tasks:
            if not json_mode:
                print(f"YAML ブロックを持つスライドが見つかりません: {pptx_path}")
            json_entries.append({"pptx": str(pptx_path), "tasks": [], "writes": []})
            continue

        if not json_mode:
            print(f"対象スライド ({len(tasks)} 件): {pptx_path}")
            print(f"{'スライド':>6}  {'id':<18}  {'baseframe':<18}  {'slots':>5}  出力先")
            print("-" * 100)
            for task in tasks:
                print(
                    f"{task['slide']:>6}  {task['id']:<18}  {task['baseframe']:<18}  "
                    f"{task['slots']:>5}  {task['output']}"
                )

        write_results: list[dict] = []
        for task in tasks:
            slide_n: int = task["slide"]
            out_path: Path = task["output"]

            if out_path.exists() and not force:
                logging.error("出力ファイルが既に存在します (--force で上書き): %s", out_path)
                errors += 1
                failures.append(f"slide {slide_n}: 出力が既存で --force なし: {out_path}")
                continue

            if dry_run:
                emit(args, "ok", "pptyamlexport", pptx_path, out_path, dry_run=True)
                write_results.append({"slide": slide_n, "id": task["id"], "output": str(out_path)})
                continue

            try:
                out_path.parent.mkdir(parents=True, exist_ok=True)
                with out_path.open("w", encoding="utf-8", newline="\n") as f:
                    f.write(_dump_doc(task["doc"]))
                emit(args, "ok", "pptyamlexport", pptx_path, out_path)
                write_results.append({"slide": slide_n, "id": task["id"], "output": str(out_path)})
            except Exception as e:
                logging.error("書き込み失敗: slide %d: %s", slide_n, e)
                errors += 1
                failures.append(f"slide {slide_n}: 書き込み失敗: {out_path}: {e}")

        if getattr(args, "closepptx", False):
            del prs

        json_entries.append({
            "pptx": str(pptx_path),
            "tasks": [
                {k: (str(v) if k == "output" else v) for k, v in t.items() if k != "doc"}
                for t in tasks
            ],
            "writes": write_results,
        })

    if json_mode:
        print(json.dumps({
            "status": "ok" if errors == 0 else "error",
            "command": "pptyamlexport",
            "entries": json_entries,
            "failures": failures,
        }, ensure_ascii=False, indent=2))
    elif failures:
        # 個々の失敗は上に出ていても OK 行に埋もれやすいので終了直前に再掲する
        # （pptimgexport と同じ方針）。
        logging.error("pptyamlexport: %d 件のタスクが失敗しました:", len(failures))
        for reason in failures:
            logging.error("  - %s", reason)

    return EXIT_OK if errors == 0 else EXIT_FAILURE
