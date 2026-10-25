"""CSL マーカー（``#CSL#``）専用の処理。

マーカーを次の 2 種類に分けて扱う。

* **一般マーカー** — ``#CSL#`` 以外のすべて（``#CSP#`` / ``#TEMP#`` / ``#MEMO#`` /
  ``#CHAPT#`` / ``#SECTION#`` / ``#STAY#`` など）。個別 pptx の段階で処理してよい。
* **CSL マーカー** — ``#CSL#`` のみ。本モジュールが扱う。

CSL マーカーへの処理は 2 種類ある。

* **CSL スライド削除処理** (``delete``) — ``#CSL#`` を含むスライドを削除する
* **CSL マーカー消去処理** (``strip``) — スライドは残しマーカー文字列のみ消す

全体方針として **CSL スライド削除は pptmerge の最後（結合済み pptx）でのみ行う**。
個別 pptx の前処理や ppt_to_* パイプラインでは ``keep``（未処理）とする。

win32com を使わず python-pptx のみで動作する。
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from ._slide_operation import slide_all_text

CSL_MARKER = "#CSL#"

#: CSL マーカーの処理モード。
#: ``delete`` = CSL スライド削除処理 / ``strip`` = CSL マーカー消去処理 / ``keep`` = 未処理
CSL_MODES = ("delete", "strip", "keep")

# ラン内の #CSL# 以降を行末まで落とす。ランをまたいだ文字は残る（元実装と同じ制約）。
_CSL_STRIP_RE = re.compile(r"#CSL#\s?.*")


def resolve_csl_mode(deletecsl: Any) -> str:
    """レシピの ``indexer.deletecsl`` を CSL 処理モードへ写す。

    ``True`` → ``delete``（CSL スライド削除処理）/ ``False`` → ``strip``（CSL マーカー消去処理）。
    ``"keep"`` / ``"delete"`` / ``"strip"`` の文字列を直接書いた場合はそれを優先する。
    """
    if isinstance(deletecsl, str) and deletecsl.strip().lower() in CSL_MODES:
        return deletecsl.strip().lower()
    return "delete" if deletecsl else "strip"


def iter_text_frames(shape: Any) -> Iterator[Any]:
    """*shape* 配下のテキストフレームを列挙する（グループ・表セルを含む）。"""
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            yield from iter_text_frames(child)
        return
    if getattr(shape, "has_table", False) and shape.table is not None:
        for row in shape.table.rows:
            for cell in row.cells:
                if cell.text_frame is not None:
                    yield cell.text_frame
        return
    if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
        yield shape.text_frame


def strip_marker_in_run(run: Any) -> bool:
    """ラン内の ``#CSL#`` 以降を除去する。除去したら ``True``。"""
    if CSL_MARKER not in run.text:
        return False
    run.text = _CSL_STRIP_RE.sub("", run.text)
    return True


def find_marked_slides(prs: Any, marker: str = CSL_MARKER) -> list[int]:
    """*marker* を含むスライドの 1 始まり番号を返す（グループ・表も走査）。"""
    return [
        i for i, slide in enumerate(prs.slides, start=1)
        if marker in slide_all_text(slide)
    ]


def delete_slides(prs: Any, indices: list[int]) -> None:
    """1 始まりの *indices* のスライドを削除し、パート関連（rel）も破棄する。"""
    id_list = prs.slides._sldIdLst
    entries = list(id_list)
    for i in sorted(indices, reverse=True):
        sld_id = entries[i - 1]
        rel_id = sld_id.get(qn("r:id"))
        if rel_id:
            prs.part.drop_rel(rel_id)
        id_list.remove(sld_id)
