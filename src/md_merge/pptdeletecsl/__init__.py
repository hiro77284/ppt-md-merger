"""pptdeletecsl — 既存 pptx に対して ``#CSL#`` の処理だけを単独で行う。

``pptmerge`` の CSL マーカー処理を、任意の既存 pptx に後から適用するためのサブコマンド。
入力はレシピ YAML ではなく pptx ファイルを CLI で直接指定する。

処理の実体は ``md_merge.pptmerge._csl`` を共有しており、``pptmerge`` の postprocess が
結合済み pptx に対して行う処理と同一。PowerPoint COM は使わない。

モード:
  ``delete`` — ``#CSL#`` を含むスライドを丸ごと削除する（既定）
  ``strip``  — スライドは残し、``#CSL#`` 以降のマーカー文字列だけを除去する
"""

from __future__ import annotations

import argparse
import logging
from pathlib import Path
from typing import Any

from pptx import Presentation

from md_merge._output import EXIT_BAD_INPUT, EXIT_FAILURE, EXIT_NOT_FOUND, EXIT_OK, emit, setup_logging
from md_merge.pptmerge._csl import (
    CSL_MARKER,
    delete_slides,
    find_marked_slides,
    iter_text_frames,
    strip_marker_in_run,
)
from md_merge.pptmerge._slide_operation import slide_all_text
from md_merge.pptmerge._slide_titles import NOREPLACE_MARKER


def strip_markers(prs: Any, marker: str = CSL_MARKER) -> tuple[int, list[int]]:
    """Strip ``#CSL#...`` text in place. Returns (runs modified, slides touched).

    ``#NOREPLACE#`` スライドは pptmerge の modify_slide_shapes と同様に対象外。
    """
    modified_runs = 0
    touched: list[int] = []
    for i, slide in enumerate(prs.slides, start=1):
        if NOREPLACE_MARKER in slide_all_text(slide):
            logging.debug("pptdeletecsl: スライド %d は %s のためスキップ", i, NOREPLACE_MARKER)
            continue
        slide_hits = 0
        for shape in slide.shapes:
            for text_frame in iter_text_frames(shape):
                for paragraph in text_frame.paragraphs:
                    for run in paragraph.runs:
                        if marker in run.text and strip_marker_in_run(run):
                            slide_hits += 1
        if slide_hits:
            modified_runs += slide_hits
            touched.append(i)
    return modified_runs, touched


def _resolve_paths(args: argparse.Namespace) -> tuple[Path, Path] | int:
    raw_input = getattr(args, "input", None)
    if not raw_input:
        logging.error("処理対象の pptx が指定されていません")
        return EXIT_BAD_INPUT

    in_path = Path(raw_input).resolve()
    if not in_path.exists():
        logging.error("入力ファイルが見つかりません: %s", in_path)
        return EXIT_NOT_FOUND
    if in_path.suffix.lower() != ".pptx":
        logging.error("入力は .pptx ファイルを指定してください: %s", in_path)
        return EXIT_BAD_INPUT

    out_raw = getattr(args, "output", None)
    out_path = Path(out_raw).resolve() if out_raw else in_path
    if out_path != in_path and out_path.exists() and not getattr(args, "force", False):
        logging.error("出力ファイルが既に存在します（--force で上書き）: %s", out_path)
        return EXIT_FAILURE
    return in_path, out_path


def run(args: argparse.Namespace) -> int:
    setup_logging(getattr(args, "log_level", "info"))

    resolved = _resolve_paths(args)
    if isinstance(resolved, int):
        return resolved
    in_path, out_path = resolved

    mode = getattr(args, "mode", "delete")
    in_place = out_path == in_path

    try:
        prs = Presentation(str(in_path))
    except Exception as e:  # noqa: BLE001 - 破損 pptx 等をそのまま失敗として扱う
        logging.error("pptx を開けませんでした: %s: %s", in_path, e)
        return EXIT_FAILURE

    total = len(prs.slides._sldIdLst)

    if mode == "delete":
        targets = find_marked_slides(prs)
        if targets and len(targets) == total:
            logging.error(
                "全 %d スライドが %s を含むため、削除すると空の pptx になります: %s",
                total, CSL_MARKER, in_path,
            )
            return EXIT_FAILURE
        summary = (
            f"{CSL_MARKER} を含むスライド {len(targets)}/{total} 枚を削除"
            f"{'（対象なし）' if not targets else ': ' + ', '.join(map(str, targets))}"
        )
    else:
        targets = find_marked_slides(prs)
        summary = f"{CSL_MARKER} マーカーを除去（対象スライド {len(targets)}/{total} 枚）"

    if getattr(args, "dry_run", False):
        logging.info("%s", summary)
        if not getattr(args, "json", False):
            print(summary)
        emit(args, "ok", "pptdeletecsl", in_path, out_path, dry_run=True)
        return EXIT_OK

    try:
        if mode == "delete":
            delete_slides(prs, targets)
            logging.info("pptdeletecsl: %d スライドを削除しました", len(targets))
        else:
            runs, touched = strip_markers(prs)
            summary = (
                f"{CSL_MARKER} マーカーを {runs} 箇所除去"
                f"（スライド {len(touched)}/{total} 枚）"
            )
            logging.info("pptdeletecsl: %d 箇所のマーカーを除去しました", runs)
        prs.save(str(out_path))
    except Exception as e:  # noqa: BLE001 - 保存失敗を含めステップ失敗として返す
        logging.error("pptdeletecsl に失敗しました: %s: %s", in_path, e)
        return EXIT_FAILURE

    if not getattr(args, "json", False):
        print(summary)
        if in_place:
            print(f"（入力ファイルを直接更新しました: {out_path}）")

    emit(args, "ok", "pptdeletecsl", in_path, out_path)
    return EXIT_OK
